"""Read a supplied budget workbook; create a private, repeat-safe SQL import.

Run with the bundled Python/openpyxl. Outputs MUST stay outside public/ and git.
The source workbook is never modified. No workbook data is embedded in this script.
"""
import datetime
import json
import pathlib
import sys
from decimal import Decimal
import openpyxl

source = pathlib.Path(sys.argv[1]).resolve()
destination = pathlib.Path(sys.argv[2]).resolve()
repo = pathlib.Path(__file__).resolve().parents[1]
if not destination.is_relative_to(repo / 'outputs'):
    raise SystemExit('Use the ignored outputs directory for private financial records.')
destination.mkdir(parents=True, exist_ok=True)
book = openpyxl.load_workbook(source, data_only=True)
transactions = []
def text(v): return '' if v is None else str(v)
def flag(v):
    if v in (None, False, 'FALSE', 'False', 0): return False
    if v in (True, 'TRUE', 'True', 1): return True
    raise ValueError('Unexpected boolean value')

for sheet, kind, amount_column in [('Revenue','revenue',5),('Purchases','expense',6)]:
    for row_number, row in enumerate(book[sheet].iter_rows(min_row=2, values_only=True),2):
        if all(v is None for v in row): continue
        if not isinstance(row[0], (datetime.datetime, datetime.date)) or row[amount_column] is None:
            raise ValueError(f'Review {sheet} row {row_number}: incomplete date or amount.')
        amount = Decimal(str(row[amount_column]))
        if amount < 0 or amount != amount.quantize(Decimal('0.01')): raise ValueError('Invalid amount')
        item = dict(kind=kind, occurred_on=row[0].strftime('%Y-%m-%d'), category=text(row[1]),
                    counterparty=text(row[2]),description=text(row[3]),program=text(row[4]),
                    payment_method=text(row[5]) if kind=='expense' else '',amount=float(amount),
                    restricted=flag(row[6]) if kind=='revenue' else False,
                    reimbursable=flag(row[7]) if kind=='expense' else False,
                    paid=flag(row[8]) if kind=='expense' else False,
                    status=text(row[7]) if kind=='revenue' else 'Received',
                    notes=text(row[8] if kind=='revenue' else row[9]),receipt_url='',grant_id=None,
                    source_key=f'fom-operating-budget-v1:{sheet}:{row_number}',source_label=f'{source.name} / {sheet} row {row_number}')
        item['in_kind'] = item['category']=='In-Kind Value' or item['payment_method']=='In-Kind'
        transactions.append(item)
years = {int(t['occurred_on'][:4]) for t in transactions}
if len(years)!=1: raise ValueError('Multiple years: confirm which year owns the undated budget before importing.')
year = next(iter(years))
budgets=[]
for kind, col, start, stop in [('revenue',1,4,12),('expense',4,4,17)]:
    for r in range(start,stop+1):
        category=book['Budget Setup'].cell(r,col).value
        value=book['Budget Setup'].cell(r,col+1).value
        if category is None or value is None: raise ValueError('Missing budget category or amount')
        budgets.append(dict(year=year,kind=kind,category=category,amount=float(Decimal(str(value))),notes=f'Imported from {source.name}. Calendar year inferred from the dated transactions.'))

def literal(v):
    if v is None: return 'null'
    if isinstance(v,bool): return 'true' if v else 'false'
    if isinstance(v,(int,float)): return str(v)
    return "'"+str(v).replace("'","''")+"'"

schema=(repo/'supabase/finance-hub-setup.sql').read_text()
statements=[schema,'\n-- PRIVATE WORKBOOK IMPORT. Do not publish or commit this file.\nbegin;']
for table,items,conflict in [('finance_transactions',transactions,'source_key'),('finance_budgets',budgets,'year,kind,category')]:
    for item in items:
        statements.append(f"insert into public.{table} ({','.join(item)}) values ({','.join(literal(v) for v in item.values())}) on conflict ({conflict}) do nothing;")
statements.append('commit;')
statements.append("select kind,count(*) as imported_records,sum(amount) as imported_amount from public.finance_transactions where source_key like 'fom-operating-budget-v1:%' group by kind order by kind;")
(destination/'activate-finance-hub.sql').write_text('\n'.join(statements))
(destination/'workbook-import.json').write_text(json.dumps({'transactions':transactions,'budgets':budgets},indent=2))
summary={kind:{'count':sum(t['kind']==kind for t in transactions),'total':str(sum((Decimal(str(t['amount'])) for t in transactions if t['kind']==kind),Decimal(0)))} for kind in ('revenue','expense')}
summary['review'] = {'Revenue!A1':book['Revenue']['A1'].value,'action':'Preserved in original workbook; excluded from transaction import because it is a header cell.'}
(destination/'reconciliation.json').write_text(json.dumps(summary,indent=2))
print(json.dumps({'records':len(transactions),'budget_lines':len(budgets),'year':year,'output':str(destination/'activate-finance-hub.sql')}))
